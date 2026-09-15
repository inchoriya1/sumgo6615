# -*- coding: utf-8 -*-
u"""
보이는 것 두 가지를 고칩니다 — 요일 색과 근퇴계 출퇴근 시각.

  python tools/fix-display.py 급여대장.xlsx            # 보기만
  python tools/fix-display.py 급여대장.xlsx --write    # 새 파일 만들기
  python tools/fix-display.py 급여대장.xlsx --write --out 같은이름.xlsx   # 덮어쓰기

① 토요일은 파랑, 일요일·공휴일은 빨강

  달력과 같은 규칙입니다. **근태가 없는 날도 칠합니다** — 근태를 넣을 때
  무슨 날인지 한눈에 들어오라고 그렇게 했습니다.

  · 파랑 — 그냥 토요일. 그리고 **대체공휴일이 딸린 토요일 공휴일**
  · 빨강 — 일요일. 그리고 **대체공휴일이 없는 공휴일**

  **토요일 공휴일이 갈립니다.** 대체공휴일이 따로 생기면 그날은 쉬는 날이
  뒤로 넘어가므로 **토요일 취급(파랑)**, 안 생기면 그날이 곧 공휴일이라
  **빨강**입니다. 2026년으로 보면 —

      8/15 광복절(토)   → 8/17 대체공휴일 있음  → 파랑
      10/3 개천절(토)   → 10/5 대체공휴일 있음  → 파랑
      6/6  현충일(토)   → 대체 없음            → 빨강
      9/26 추석연휴(토) → 대체 없음            → 빨강

  **표를 그대로 읽습니다** — 어느 공휴일에 대체가 붙었는지를 박아 넣지 않고,
  [월설정] G·H 열에서 **그 날 뒤 사흘 안에 '대체공휴일' 이 있는지**를 셉니다.
  토요일 공휴일의 대체는 다음 월요일(+2)이라 사흘이면 넉넉합니다.
  그래서 **해가 바뀌어 표만 채워 넣으면 색도 따라갑니다.**

  글자색만 바꿉니다. 칸을 칠하면 이미 있는 **사람 블록 줄무늬**(연파랑)와
  **노란 입력칸**을 덮어 버려서, 넣을 자리를 못 찾게 됩니다.

② 색만 파랗고 구분은 '특근' 이던 것

  `U` 열(구분)의 토요 판정이 `$C2="휴일"` 만 보고 있었습니다.
  그래서 **광복절처럼 이름이 붙은 날은 토요일이어도 무조건 '특근'** 으로 갔습니다.
  색은 파란데 구분은 특근인, 말이 안 맞는 상태였습니다.

  **색과 똑같은 조건**을 쓰게 고칩니다 — `sub_expr()` 하나를 양쪽이 나눠 씁니다.

      전  IF(AND($R2="토",$C2="휴일"),"토요","특근")
      후  IF(AND($R2="토",OR($C2="휴일",대체있음)),"토요","특근")

  **급여가 바뀝니다.** 특근과 토요는 계산이 다릅니다 —
  · 휴일특근수당(AB) = 통상시급 × 1.5 × **실적시간**(근태집계 H)
  · 토요근무수당(Z)  = 통상시급 × 1.5 × **고정시간**(직원설정/월별입력)
    거기에 실적이 고정에 못 미치면 **토요근무차감(AC)** 이 붙습니다.

  그래서 8H 가 특근에서 토요로 옮겨가면 **특근수당 8H 어치가 빠지고**,
  그만큼 **토요 실적이 채워져 차감이 줄어듭니다.** 차감이 이미 0이던 사람은
  줄어들 차감이 없어 **그냥 빠지기만 합니다.** 사람마다 다르니 확인이 필요합니다.

③ 근퇴계 출근·퇴근이 `0.354166666666667` 로 보입니다

  `/24` 를 하는 수식은 **하나도 없었습니다.** 원인은 다른 것이었습니다 —
  근퇴계가 근태원본에서 시각을 끌어오며 끝에 `&""` 를 붙이고 있었습니다.

      =IFERROR(INDEX(근태원본!$D:$D, MATCH(...))&"", "")
                                              ~~~~~

  `&""` 는 **숫자를 글자로 바꿉니다.** 엑셀에서 시각은 '하루의 몇 분의 몇'인
  숫자(8:30 = 0.354166…)라, 글자가 되는 순간 **서식이 듣지 않습니다.**
  `h:mm` 을 입혀도 소용없습니다. 이미 글자니까요.

  빈 날에 `0` 이 뜨는 것을 막으려고 붙인 것으로 보입니다. 그래서 `&""` 를
  떼고, 대신 **서식의 0 자리를 비워**(`h:mm;h:mm;`) 같은 효과를 냅니다.

  근무일명칭(C)·비고(K)도 `&""` 를 쓰지만 **그쪽은 원래 글자**라 그대로 둡니다.
"""
import argparse
import re

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font

RAW = u"근태원본"
GT = u"근퇴계"

RAW_FIRST, RAW_LAST = 2, 5953
GT_FIRST, GT_LAST = 8, 38

RED = u"FFC00000"       # 일요일 · 대체공휴일이 없는 공휴일
BLUE = u"FF0070C0"      # 토요일 · 대체공휴일이 딸린 토요일 공휴일

SUB_WINDOW = 3          # 공휴일 뒤 며칠 안의 '대체공휴일' 을 그 공휴일 것으로 볼지

TIME_FMT = u"h:mm;h:mm;"        # 셋째 칸(0)을 비워 빈 날을 안 보이게 합니다


def sub_expr(date_col, row):
    u"""그 공휴일에 대체공휴일이 딸려 있나 — 색과 구분이 **똑같이** 쓰는 식입니다.

    토요일 공휴일의 대체는 다음 월요일(+2), 일요일 공휴일은 다음 날(+1)이라
    사흘이면 넉넉히 덮습니다. 표를 그대로 읽으므로 해가 바뀌어도 따라갑니다.
    """
    return (u'COUNTIFS(공휴일날짜,">"&${d}{r},공휴일날짜,"<="&${d}{r}+{w},'
            u'공휴일이름,"대체공휴일")>0').format(d=date_col, r=row, w=SUB_WINDOW)


# U 구분 — 색과 같은 조건을 씁니다. 안 그러면 '파란데 특근' 이 생깁니다.
CLASS_FORMULA = (
    u'=IF(N($T{r})=0,"",IF($C{r}="평일","연장",'
    u'IF(AND($R{r}="토",OR($C{r}="휴일",{sub})),"토요","특근")))'
)


def fix_class(wb):
    u"""근태원본 U열(구분) — 대체공휴일이 딸린 토요일 공휴일을 '토요' 로 잡습니다."""
    ws = wb[RAW]
    before = ws["U2"].value
    for r in range(RAW_FIRST, RAW_LAST + 1):
        ws.cell(row=r, column=21).value = CLASS_FORMULA.format(
            r=r, sub=sub_expr("B", r))
    return before, ws["U2"].value


def day_rules(date_col, dow_col, name_col, first):
    u"""(빨강 규칙, 파랑 규칙). 열 문자만 시트마다 갈아 끼웁니다.

    수식 앞에 `=` 를 붙이지 않습니다 — 조건부서식 XML 은 `=` 없이 담고,
    붙이면 엑셀이 `==AND(...)` 로 읽어 규칙이 조용히 죽습니다.
    이미 있는 줄무늬 규칙도 `=` 없이 들어 있습니다.
    """
    d, w, n = date_col, dow_col, name_col

    sub = sub_expr(d, first)

    is_hol = (u'AND(${n}{r}<>"",${n}{r}<>"평일",${n}{r}<>"휴일")'
              ).format(n=n, r=first)

    # 파랑 — 그냥 토요일이거나, 대체공휴일이 딸린 토요일 공휴일
    blue = (u'AND(${d}{r}<>"",${w}{r}="토",OR(${n}{r}="휴일",{sub}))'
            ).format(d=d, w=w, n=n, r=first, sub=sub)

    # 빨강 — 일요일이거나, 공휴일이면서 위 파랑에 해당하지 않는 날
    red = (u'AND(${d}{r}<>"",OR(${w}{r}="일",'
           u'AND({hol},NOT(AND(${w}{r}="토",{sub})))))'
           ).format(d=d, w=w, r=first, hol=is_hol, sub=sub)
    return (FormulaRule(formula=[red], font=Font(color=RED)),
            FormulaRule(formula=[blue], font=Font(color=BLUE)))


def is_day_rule(rule):
    u"""전에 이 도구가 넣은 요일 색 규칙인가 — 두 번 돌려도 겹치지 않게 골라냅니다."""
    dxf = rule.dxf
    if not (dxf and dxf.font and dxf.font.color):
        return False
    return getattr(dxf.font.color, "rgb", None) in (RED, BLUE)


def color_days(wb):
    done = []

    # 근태원본 — B 근무일자 · R 요일 · C 근무일명칭
    ws = wb[RAW]
    keep = [rule for cf in ws.conditional_formatting for rule in cf.rules
            if not is_day_rule(rule)]
    ref = u"A%d:Y%d" % (RAW_FIRST, RAW_LAST)
    # 이름(A열)은 요일 색에서 뺍니다 — 사람 이름이 빨강·파랑으로 물들면
    # 무슨 표시인가 싶어 눈이 갑니다. 줄무늬는 A열까지 그대로 둡니다.
    color_ref = u"B%d:Y%d" % (RAW_FIRST, RAW_LAST)
    ws.conditional_formatting = ConditionalFormattingList()
    for rule in keep:                       # 줄무늬·블록 경계선을 먼저 둡니다
        ws.conditional_formatting.add(ref, rule)
    for rule in day_rules("B", "R", "C", RAW_FIRST):
        ws.conditional_formatting.add(color_ref, rule)
    done.append((RAW, color_ref, len(keep), 2))

    # 근퇴계 — A 일자 · B 요일 · C 근무일명칭
    ws = wb[GT]
    keep = [rule for cf in ws.conditional_formatting for rule in cf.rules
            if not is_day_rule(rule)]
    ref = u"A%d:K%d" % (GT_FIRST, GT_LAST)
    ws.conditional_formatting = ConditionalFormattingList()
    for rule in keep:
        ws.conditional_formatting.add(ref, rule)
    for rule in day_rules("A", "B", "C", GT_FIRST):
        ws.conditional_formatting.add(ref, rule)
    done.append((GT, ref, len(keep), 2))

    return done


AMP = re.compile(r'\)&""')


def fix_times(wb):
    u"""근퇴계 D·E 의 `&\"\"` 를 떼고 시각 서식을 입힙니다."""
    ws = wb[GT]
    fixed = []
    for col in ("D", "E"):
        for r in range(GT_FIRST, GT_LAST + 1):
            c = ws["%s%d" % (col, r)]
            v = c.value
            if isinstance(v, str) and '&""' in v:
                c.value = AMP.sub(")", v)
                fixed.append(c.coordinate)
            c.number_format = TIME_FMT
    return fixed


def run(path, out_path, do_write):
    wb = openpyxl.load_workbook(path)

    before = wb[GT]["D8"].value
    done = color_days(wb)
    cls_before, cls_after = fix_class(wb)
    fixed = fix_times(wb)

    print(u"① 토요일 파랑 · 일요일·공휴일 빨강 (글자색)")
    for sheet, ref, keep, added in done:
        print(u"   %-6s %-12s 원래 규칙 %d개 유지 + %d개 추가"
              % (sheet, ref, keep, added))
    print(u"② 근태원본 U열 구분 — %d줄. 색과 같은 조건을 씁니다"
          % (RAW_LAST - RAW_FIRST + 1))
    print(u"   전: %s" % cls_before)
    print(u"   후: %s" % cls_after[:120])
    print(u"③ 근퇴계 출근·퇴근 — %d칸에서 `&\"\"` 를 떼고 서식 %s"
          % (len(fixed), TIME_FMT))
    print(u"   전: %s" % before[:96])
    print(u"   후: %s" % wb[GT]["D8"].value[:96])

    if not do_write:
        print(u"\n※ 보기만 했습니다. 실제로 만들려면 --write 를 붙이세요.")
        return
    wb.save(out_path)
    print(u"\n만들었습니다: %s" % out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("book")
    ap.add_argument("--out", default=None)
    ap.add_argument("--write", action="store_true")
    a = ap.parse_args()
    out = a.out or re.sub(r"\.xlsx$", "", a.book) + u"_보기.xlsx"
    run(a.book, out, a.write)


if __name__ == "__main__":
    main()
