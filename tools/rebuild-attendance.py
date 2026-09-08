# -*- coding: utf-8 -*-
"""
근태원본을 '출퇴근만 넣으면 되는' 구조로 바꿉니다.

  python tools/rebuild-attendance.py 급여대장.xlsx            # 보기만
  python tools/rebuild-attendance.py 급여대장.xlsx --write    # 새 파일 만들기

지금까지
  근태기록 .xls 의 **열네 개 열**을 통째로 복사해 맨 아래 빈 줄에 붙여넣었습니다.
  자리를 찾아야 하고, 머리글·합계 줄을 빼야 하고, 두 번 넣으면 조용히 두 배가 됐습니다.

바꾼 뒤
  연월 × 사람 × 31일 자리를 **미리 깔아 둡니다**(8명 × 24개월 × 31일 = 5,952줄).
  이름 · 날짜 · 요일 · 근무일명칭은 **저절로** 채워집니다.
  사람이 넣는 것은 **출근·퇴근 두 열뿐**입니다.
  붙여넣을 자리도 늘 정해져 있어서, 근태넣기가 `D1800:E1830` 처럼 딱 집어 줍니다.

기본·연장 시간을 어떻게 내나
  이 회사의 근태기록 **475줄에서 규칙을 찾아내어 대조**했습니다(472줄 일치).

      평일   기본 8시간,  연장 = (퇴근 − 17:30) 을 30분 단위로 내림
      휴일   기본 0,      연장 = (퇴근 − 08:30) 을 30분 단위로 내림

  맞지 않던 3줄은 전부 **사람이 판단해 손으로 깎은 지각**이었습니다.
  그래서 O열에 **지각차감(시간)** 칸을 두었습니다. 거기에 적으면 기본에서 뺍니다.

공휴일
  근무일명칭은 **요일 + 공휴일 표**로 정해집니다.
  표는 [월설정] 시트 G열에 있습니다. **비어 있으면 그날이 평일로 잡히니**
  해마다 채워 주세요. 근태넣기가 그 달 공휴일 수를 보여 줍니다.
"""
import argparse
import datetime
import re

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

RAW, SET_EMP, SET_MON = u"근태원본", u"직원설정", u"월설정"
AGG, GT, IN = u"근태집계", u"근퇴계", u"근태넣기"
SLOTS, MONTHS, DAYS = 8, 24, 31
FIRST = 2
LAST = FIRST + SLOTS * MONTHS * DAYS - 1          # 5953
OLD_LAST = 4001
EMP_ROW0, MON_ROW0 = 2, 4                          # 직원설정 B2.., 월설정 A4..
HOL_COL, HOL_ROW0, HOL_N = 7, 4, 60                # 월설정 G4:H63

FONT = u"맑은 고딕"
RED = u"FFC00000"
FILL_IN = PatternFill("solid", fgColor=u"FFFFE699")
FILL_AUTO = PatternFill("solid", fgColor=u"FFF2F2F2")
FILL_HEAD = PatternFill("solid", fgColor=u"FF1F4E79")
FILL_HEAD2 = PatternFill("solid", fgColor=u"FFC55A11")
THIN = Side(style="thin", color=u"FFD9D9D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# 자료에서 찾아낸 공휴일. 확인된 것만 넣습니다 — 나머지는 사용자가 채웁니다.
KNOWN_HOLIDAYS = [
    (datetime.date(2026, 6, 3), u"지방선거일"),
    (datetime.date(2026, 6, 6), u"현충일"),
    (datetime.date(2026, 7, 17), u"제헌절"),
    (datetime.date(2026, 8, 15), u"광복절"),
    (datetime.date(2026, 8, 17), u"대체공휴일"),
]


def block_start(month_i, slot_i):
    """month_i·slot_i 는 0부터. 그 사람 그 달의 첫 줄 번호."""
    return FIRST + (month_i * SLOTS + slot_i) * DAYS


def widen_refs(wb, report):
    """근태원본 범위 $4001 → $5953 로 넓힙니다."""
    pat = re.compile(r"(%s!\$[A-Z]{1,2}\$%d:\$[A-Z]{1,2}\$)%d" % (RAW, FIRST, OLD_LAST))
    n = 0
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for c in row:
                if isinstance(c.value, str) and u"$%d" % OLD_LAST in c.value:
                    new = pat.sub(lambda m: m.group(1) + str(LAST), c.value)
                    if new != c.value:
                        c.value = new
                        n += 1
    report.append(u"근태원본 참조 범위를 %d → %d 로 넓혔습니다 (%d칸)" % (OLD_LAST, LAST, n))


def read_existing(wb):
    """지금 들어 있는 근태를 (이름, 날짜) → 값 으로 거둬들입니다."""
    raw = wb[RAW]
    out = {}
    for r in range(FIRST, OLD_LAST + 1):
        nm = raw.cell(r, 1).value
        if not nm:
            continue
        d = raw.cell(r, 2).value
        key = (unicode_(nm), date_key(d))
        # 평일인데 기본이 8시간이 아니면, 사람이 손으로 깎은 지각입니다.
        # 그 차이를 O열(지각차감)로 옮겨 담아야 시간이 그대로 재현됩니다.
        late = 0
        if raw.cell(r, 3).value == u"평일":
            hrs = _time_hours(raw.cell(r, 11).value)
            if hrs is not None and abs(hrs - 8) > 1e-6:
                late = round(8 - hrs, 2)
        out[key] = dict(ins=raw.cell(r, 4).value, out=raw.cell(r, 5).value,
                        memo=raw.cell(r, 14).value, late=late)
    return out


def _time_hours(v):
    """엑셀이 시간을 담는 방식이 여럿입니다 — timedelta · time · 숫자 · 글자."""
    if v is None:
        return None
    if isinstance(v, datetime.timedelta):
        return v.total_seconds() / 3600.0
    if isinstance(v, (int, float)):
        return v * 24.0
    if hasattr(v, "hour"):
        return v.hour + v.minute / 60.0 + v.second / 3600.0
    m = re.match(r"^\s*(\d+):(\d+)(?::(\d+))?\s*$", unicode_(v))
    if m:
        h, mi, se = (int(x or 0) for x in m.groups())
        return h + mi / 60.0 + se / 3600.0
    return None


def unicode_(v):
    return v if isinstance(v, str) else str(v)


def date_key(v):
    if v is None:
        return None
    if hasattr(v, "year"):
        return u"%04d-%02d-%02d" % (v.year, v.month, v.day)
    s = unicode_(v).strip().replace("/", "-")
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    return u"%04d-%02d-%02d" % tuple(int(x) for x in m.groups()) if m else None


def holiday_table(wb, report):
    ms = wb[SET_MON]
    put(ms, "G3", u"공휴일 날짜", 9, True, u"FFFFFFFF", FILL_HEAD, True, "center")
    put(ms, "H3", u"이름", 9, True, u"FFFFFFFF", FILL_HEAD, True, "center")
    put(ms, "I3", u"← 여기가 비면 그날이 '평일'로 잡힙니다. 해마다 채워 주세요",
        10, color=RED)
    for i in range(HOL_N):
        r = HOL_ROW0 + i
        for c in (HOL_COL, HOL_COL + 1):
            cell = ms.cell(r, c)
            cell.fill = FILL_IN
            cell.border = BOX
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(horizontal="center")
    for i in range(HOL_N):                      # 날짜 서식을 미리 입혀 둡니다
        ms.cell(HOL_ROW0 + i, HOL_COL).number_format = "yyyy-mm-dd"
    for i, (d, nm) in enumerate(KNOWN_HOLIDAYS):
        r = HOL_ROW0 + i
        ms.cell(r, HOL_COL).value = d           # 글자가 아니라 진짜 날짜여야 MATCH 가 맞습니다
        ms.cell(r, HOL_COL + 1).value = nm
    ms.column_dimensions["G"].width = 14
    ms.column_dimensions["H"].width = 14
    ms.column_dimensions["I"].width = 46
    for nm, ref in ((u"공휴일날짜", u"%s!$G$%d:$G$%d" % (SET_MON, HOL_ROW0, HOL_ROW0 + HOL_N - 1)),
                    (u"공휴일이름", u"%s!$H$%d:$H$%d" % (SET_MON, HOL_ROW0, HOL_ROW0 + HOL_N - 1))):
        if nm in wb.defined_names:
            del wb.defined_names[nm]
        wb.defined_names[nm] = openpyxl.workbook.defined_name.DefinedName(nm, attr_text=ref)
    report.append(u"[월설정] G열에 공휴일 표를 만들고 확인된 %d개를 넣었습니다"
                  % len(KNOWN_HOLIDAYS))


def put(ws, addr, value, size=11, bold=False, color=None, fill=None,
        box=False, align=None, wrap=False):
    c = ws[addr] if isinstance(addr, str) else addr
    c.value = value
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    if fill:
        c.fill = fill
    if box:
        c.border = BOX
    if align or wrap:
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    return c


def lay_grid(wb, existing, report):
    raw = wb[RAW]
    # 기존 A~N 지우기
    for r in range(FIRST, max(OLD_LAST, LAST) + 1):
        for c in range(1, 16):
            raw.cell(r, c).value = None

    put(raw, "O1", u"지각차감(H)", 9, True, u"FFFFFFFF", FILL_HEAD2, True, "center")
    raw.column_dimensions["O"].width = 11

    filled = 0
    for mi in range(MONTHS):
        mrow = MON_ROW0 + mi
        for si in range(SLOTS):
            erow = EMP_ROW0 + si
            top = block_start(mi, si)
            for di in range(DAYS):
                r, day = top + di, di + 1
                raw.cell(r, 1).value = (
                    u'=IF(%s!$B$%d="","",%s!$B$%d)' % (SET_EMP, erow, SET_EMP, erow))
                raw.cell(r, 2).value = (
                    u'=IF($A%d="","",IF(%d>DAY(EOMONTH(%s!$B$%d,0)),"",'
                    u'DATE(YEAR(%s!$B$%d),MONTH(%s!$B$%d),%d)))'
                    % (r, day, SET_MON, mrow, SET_MON, mrow, SET_MON, mrow, day))
                raw.cell(r, 3).value = (
                    u'=IF($B%d="","",IFERROR(INDEX(공휴일이름,MATCH($B%d,공휴일날짜,0)),'
                    u'IF(WEEKDAY($B%d,2)<=5,"평일","휴일")))' % (r, r, r))
                raw.cell(r, 11).value = (
                    u'=IF(OR($B%d="",$D%d="",$E%d=""),"",'
                    u'MAX(0,IF($C%d="평일",8,0)-N($O%d))/24)' % (r, r, r, r, r))
                raw.cell(r, 12).value = (
                    u'=IF(OR($B%d="",$D%d="",$E%d=""),"",'
                    u'FLOOR(MAX(0,$E%d-IF($C%d="평일",TIME(17,30,0),TIME(8,30,0)))*48,1)/48)'
                    % (r, r, r, r, r))
                raw.cell(r, 24).value = (
                    u'=IF($A%d="","",INT((ROW()-%d)/%d))' % (r, FIRST, DAYS))
                for c in (4, 5):                       # D 출근 · E 퇴근
                    raw.cell(r, c).fill = FILL_IN
                    raw.cell(r, c).border = BOX
                for c in (1, 2, 3, 11, 12):
                    raw.cell(r, c).fill = FILL_AUTO

            key_nm = wb[SET_EMP].cell(erow, 2).value
            if not key_nm:
                continue
            mval = wb[SET_MON].cell(mrow, 2).value      # 임금산정기간 시작(날짜)
            for di in range(DAYS):
                r, day = top + di, di + 1
                if mval is None:
                    continue
                k = u"%04d-%02d-%02d" % (mval.year, mval.month, day)
                got = existing.get((unicode_(key_nm), k))
                if not got:
                    continue
                raw.cell(r, 4).value = got["ins"]
                raw.cell(r, 5).value = got["out"]
                raw.cell(r, 14).value = got["memo"]
                if got["late"]:
                    raw.cell(r, 15).value = got["late"]
                filled += 1
    report.append(u"자리를 %d줄 깔고, 기존 근태 %d줄을 옮겨 담았습니다" % (LAST - FIRST + 1, filled))

    raw["Z1"] = (
        u"■ 매월 사용법  (이 시트는 지우지 않습니다)\n"
        u"① [근태넣기] 에서 연월과 사람을 고릅니다.\n"
        u"② 알려 주는 자리(예: D1800:E1830)를 누릅니다.\n"
        u"③ 근태기록 .xls 의 **출근·퇴근 두 열만** 복사해 붙여넣습니다.\n"
        u"   (머리글과 합계 줄은 빼고, 1일부터 차례로 맞춰서)\n\n"
        u"· 이름·날짜·요일·근무일명칭은 저절로 채워집니다. 손대지 마세요.\n"
        u"· 기본(K)·연장(L)도 퇴근 시각에서 저절로 나옵니다.\n"
        u"    평일  기본 8H, 연장 = (퇴근 − 17:30) 30분 단위 내림\n"
        u"    휴일  기본 0,  연장 = (퇴근 − 08:30) 30분 단위 내림\n"
        u"· 지각을 깎을 때만 O열 [지각차감(H)] 에 시간을 적습니다.\n"
        u"· 공휴일은 [월설정] G열 표를 보고 판단합니다. 비면 평일로 잡힙니다.")
    return filled


def rebuild_input_sheet(wb, report):
    """근태넣기를 새 구조에 맞게 다시 씁니다."""
    if IN in wb.sheetnames:
        del wb[IN]
    ws = wb.create_sheet(IN, wb.sheetnames.index(RAW))
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJ", (20, 17, 13, 13, 13, 13, 13, 13, 13, 13)):
        ws.column_dimensions[col].width = w

    cur = wb[SET_MON]["B1"].value
    first_emp = next((wb[SET_EMP].cell(EMP_ROW0 + i, 2).value for i in range(SLOTS)
                      if wb[SET_EMP].cell(EMP_ROW0 + i, 2).value), None)

    put(ws, "A1", u"▶ 넣을 연월", 12, True, RED)
    put(ws, "B1", cur, 12, True, fill=FILL_IN, box=True, align="center")
    put(ws, "C1", u"← 이 두 칸만 고르면 됩니다", 10, color=RED)
    put(ws, "A2", u"▶ 넣을 사람", 12, True, RED)
    put(ws, "B2", first_emp, 12, True, fill=FILL_IN, box=True, align="center")
    for nm, cell in ((u"연월목록", "B1"), (u"직원목록", "B2")):
        dv = DataValidation(type="list", formula1=u"=%s" % nm, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(ws[cell])

    mi = u'MATCH($B$1,%s!$A$%d:$A$%d,0)' % (SET_MON, MON_ROW0, MON_ROW0 + MONTHS - 1)
    si = u'MATCH($B$2,%s!$B$%d:$B$%d,0)' % (SET_EMP, EMP_ROW0, EMP_ROW0 + SLOTS - 1)
    top = u'(%d+((%s-1)*%d+(%s-1))*%d)' % (FIRST, mi, SLOTS, si, DAYS)
    days = (u'DAY(EOMONTH(INDEX(%s!$B$%d:$B$%d,%s),0))'
            % (SET_MON, MON_ROW0, MON_ROW0 + MONTHS - 1, mi))

    put(ws, "A4", u"■ 여기에 붙여넣으세요   (출근·퇴근 두 열만)", 11, True)
    put(ws, "A5", u"붙여넣을 자리", 11, True)
    put(ws, "B5", u'=IFERROR("D"&%s&":E"&(%s+%s-1),"고르세요")' % (top, top, days),
        16, True, fill=FILL_AUTO, box=True, align="center")
    ws.merge_cells("B5:C5")
    put(ws, "A6", u"이 달은 며칠", 11, True)
    put(ws, "B6", u'=IFERROR(%s,"")' % days, 11, fill=FILL_AUTO, box=True, align="center")
    put(ws, "A7", u"지금 채워진 날", 11, True)
    put(ws, "B7", u'=COUNTIFS(%s!$V$%d:$V$%d,$B$1&"|"&$B$2,%s!$D$%d:$D$%d,"<>")'
        % (RAW, FIRST, LAST, RAW, FIRST, LAST), 11, fill=FILL_AUTO, box=True, align="center")

    way = u'IFERROR(INDEX(%s!$I$%d:$I$%d,%s),"")' % (SET_EMP, EMP_ROW0, EMP_ROW0 + SLOTS - 1, si)
    put(ws, "A8", u"판정", 11, True)
    put(ws, "B8",
        u'=IF($B$2="","사람을 고르세요",'
        u'IF(%s="고정월급","고정월급이라 근태를 넣지 않습니다",'
        u'IF($B$7=0,"아직 안 넣었습니다 — 위 자리에 출근·퇴근을 붙여넣으세요",'
        u'$B$7&"일 채워져 있습니다 (이 달 "&$B$6&"일 중). 쉰 날은 비어 있는 것이 정상입니다")))'
        % way, 11, True, fill=FILL_AUTO, box=True)
    ws.merge_cells("B8:G8")

    put(ws, "A10", u"■ 이 달 공휴일", 11, True)
    put(ws, "B10",
        u'=IFERROR(COUNTIFS(공휴일날짜,">="&INDEX(%s!$B$%d:$B$%d,%s),'
        u'공휴일날짜,"<="&INDEX(%s!$C$%d:$C$%d,%s)),0)'
        % (SET_MON, MON_ROW0, MON_ROW0 + MONTHS - 1, mi,
           SET_MON, MON_ROW0, MON_ROW0 + MONTHS - 1, mi),
        11, True, fill=FILL_AUTO, box=True, align="center")
    put(ws, "C10",
        u'=IF($B$10=0,"★ 이 달 공휴일이 하나도 등록돼 있지 않습니다. '
        u'[월설정] G열을 확인하세요","개 등록됨 — 맞는지 [월설정] G열에서 확인하세요")',
        10, color=RED)
    ws.merge_cells("C10:H10")

    put(ws, "A12", u"■ 넣은 것 한눈에 보기   (가로 사람 · 세로 연월 · 숫자는 채워진 날 수)", 11, True)
    put(ws, "A13", u"연월", 9, True, u"FFFFFFFF", FILL_HEAD, True, "center")
    for i in range(SLOTS):
        col = chr(ord("B") + i)
        put(ws, u"%s13" % col,
            u'=IF(%s!$B$%d="","",%s!$B$%d)' % (SET_EMP, EMP_ROW0 + i, SET_EMP, EMP_ROW0 + i),
            9, True, u"FFFFFFFF", FILL_HEAD, True, "center")
    for j in range(MONTHS):
        r = 14 + j
        put(ws, u"A%d" % r, u"=%s!$A$%d" % (SET_MON, MON_ROW0 + j), 10, box=True, align="center")
        for i in range(SLOTS):
            col = chr(ord("B") + i)
            put(ws, u"%s%d" % (col, r),
                u'=IF(%s$13="","",COUNTIFS(%s!$V$%d:$V$%d,$A%d&"|"&%s$13,'
                u'%s!$D$%d:$D$%d,"<>"))'
                % (col, RAW, FIRST, LAST, r, col, RAW, FIRST, LAST),
                10, box=True, align="center")
    ws.freeze_panes = "B14"
    report.append(u"[%s] 를 새 구조에 맞게 다시 만들었습니다" % IN)


def run(path, out_path, do_write):
    wb = openpyxl.load_workbook(path, data_only=False)
    raw = wb[RAW]
    have = sum(1 for r in range(FIRST, OLD_LAST + 1) if raw.cell(r, 1).value)
    print(u"■ 지금 근태원본 — 붙여넣은 %d줄 (열네 개 열을 통째로)" % have)
    print(u"■ 바꾼 뒤 — 자리 %d줄을 미리 깔아 둡니다 (%d명 × %d개월 × %d일)"
          % (SLOTS * MONTHS * DAYS, SLOTS, MONTHS, DAYS))
    print(u"   손으로 넣는 것 : **출근(D) · 퇴근(E)** 두 열")
    print(u"   저절로 되는 것 : 이름 · 날짜 · 요일 · 근무일명칭 · 기본 · 연장")
    print(u"   새로 생기는 것 : O열 지각차감(H), [월설정] G열 공휴일 표")
    if not do_write:
        print(u"\n※ 보기만 했습니다. 실제로 만들려면 --write 를 붙이세요.")
        return

    report = []
    existing = read_existing(wb)
    holiday_table(wb, report)
    lay_grid(wb, existing, report)
    widen_refs(wb, report)
    rebuild_input_sheet(wb, report)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    print(u"")
    for line in report:
        print(u"   " + line)
    print(u"\n만들었습니다: %s" % out_path)
    print(u"원본은 그대로 두었습니다: %s" % path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("book")
    ap.add_argument("--out", default=None)
    ap.add_argument("--write", action="store_true")
    a = ap.parse_args()
    out = a.out or re.sub(r"\.xlsx$", "", a.book) + u"_근태개편.xlsx"
    run(a.book, out, a.write)
