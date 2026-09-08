# -*- coding: utf-8 -*-
"""
시트 두 곳을 '골라서 쓰는' 방식으로 바꿉니다.

  python tools/consolidate-sheets.py 급여대장.xlsx            # 보기만
  python tools/consolidate-sheets.py 급여대장.xlsx --write    # 새 파일 만들기

바꾸는 것 두 가지

① [근퇴-이름] 여섯 장 → [근퇴계] 한 장
   여섯 장은 **딱 세 칸만** 달랐습니다 — 제목, 슬롯 번호, 급여방식.
   그래서 '이 름' 칸(C3)을 **고르는 칸**으로 바꾸고 나머지를 그 칸에
   딸려 가게 했습니다. 월설정에서 달을 고르는 것과 똑같은 방식입니다.
   나머지 다섯 장은 지웁니다.

② [근태넣기] 새 시트
   연월과 사람을 고르면 **어디에 붙여넣을지** 알려 주고,
   넣은 뒤 **제대로 들어갔는지** 검사합니다.
   아래쪽에 연월 × 사람 표가 있어서 **무엇이 빠졌는지 한눈에** 보입니다.

왜 근태원본은 그대로 두나
   근태원본은 **모든 사람 · 모든 달을 쌓아 두는 창고**입니다.
   사람을 고를 때마다 내용이 바뀌게 만들면 **지난 달 기록이 사라집니다.**
   (수식으로 끌어오는 순간 저장이 아니라 '보기'가 되어 버립니다.)
   매크로 없는 .xlsx 라 단추로 옮겨 담을 수도 없습니다.
   그래서 창고는 그대로 두고, **넣는 자리를 알려 주고 검사하는 창구**를
   따로 만들었습니다.
"""
import argparse
import re

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

RAW = u"근태원본"
SET_EMP = u"직원설정"
SET_MON = u"월설정"
GUIDE = u"사용안내"
NEW_SHEET = u"근태넣기"
GT_NEW = u"근퇴계"

EMP_ROWS = (2, 9)          # 직원설정 성명 B2:B9 (8자리)
MON_ROWS = (4, 27)         # 월설정 연월 A4:A27 (24개월)

# 이 파일이 쓰고 있는 색·글꼴을 그대로 따릅니다.
FONT = u"맑은 고딕"
RED = u"FFC00000"
INK_ON_DARK = u"FFFFFFFF"
FILL_HEAD = PatternFill("solid", fgColor=u"FF1F4E79")   # 표 머리글(남색)
FILL_IN = PatternFill("solid", fgColor=u"FFFFE699")     # 직접 넣는 칸(노랑)
FILL_AUTO = PatternFill("solid", fgColor=u"FFFFF2CC")   # 자동 계산(크림)
FILL_NOTE = PatternFill("solid", fgColor=u"FFFFF9E6")   # 안내 상자
THIN = Side(style="thin", color=u"FFBFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def put(ws, addr, value, size=11, bold=False, color=None, fill=None,
        box=False, align=None, wrap=False):
    c = ws[addr]
    c.value = value
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    if fill:
        c.fill = fill
    if box:
        c.border = BOX
    if align or wrap:
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    return c


def emp_list_ref():
    return u"%s!$B$%d:$B$%d" % (SET_EMP, EMP_ROWS[0], EMP_ROWS[1])


def mon_list_ref():
    return u"%s!$A$%d:$A$%d" % (SET_MON, MON_ROWS[0], MON_ROWS[1])


# ─────────────────────────────────────────────────────────────
# ① 근퇴계 통합
# ─────────────────────────────────────────────────────────────
def merge_gt(wb, report):
    sheets = [n for n in wb.sheetnames if n.startswith(u"근퇴-")]
    if not sheets:
        report.append(u"근퇴-이름 시트가 없습니다 — 이미 통합된 것 같습니다.")
        return None
    keep = sheets[0]
    ws = wb[keep]
    report.append(u"근퇴 시트 %d장 → [%s] 한 장 (%s 를 남기고 나머지 삭제)"
                  % (len(sheets), GT_NEW, keep))

    first = None
    for r in range(EMP_ROWS[0], EMP_ROWS[1] + 1):
        v = wb[SET_EMP].cell(r, 2).value
        if v:
            first = v
            break

    ws.title = GT_NEW
    # '이 름' 칸을 고르는 칸으로
    put(ws, "C3", first, size=11, bold=True, fill=FILL_IN, box=True, align="center")
    ws.add_data_validation(_dv(u"직원목록"))
    ws.data_validations.dataValidation[-1].add(ws["C3"])
    put(ws, "L3", u"← 이 칸에서 사람을 고르세요", size=10, color=RED)

    idx = u'IFERROR(MATCH($C$3,%s,0),1)' % emp_list_ref()
    ws["N1"] = u"=월시작행+%s-1" % idx
    ws["A1"] = u'=회사명&"   [  "&$C$3&"  ]   근 퇴 계"'
    ws["G4"] = u'=IFERROR(INDEX(%s!$I$2:$I$9,%s),"")' % (SET_EMP, idx)
    ws.print_area = "A1:K62"

    for n in sheets[1:]:
        del wb[n]
    report.append(u"   지운 시트: %s" % u", ".join(sheets[1:]))
    return GT_NEW


def _dv(name):
    dv = DataValidation(type="list", formula1=u"=%s" % name, allow_blank=True)
    dv.error = u"목록에 있는 것만 고를 수 있습니다."
    dv.errorTitle = u"목록에서 고르세요"
    dv.prompt = u"눌러서 목록에서 고르세요."
    dv.promptTitle = u"고르기"
    return dv


# ─────────────────────────────────────────────────────────────
# ② 근태넣기
# ─────────────────────────────────────────────────────────────
def make_input_sheet(wb, report):
    if NEW_SHEET in wb.sheetnames:
        del wb[NEW_SHEET]
    pos = wb.sheetnames.index(RAW) if RAW in wb.sheetnames else 1
    ws = wb.create_sheet(NEW_SHEET, pos)
    ws.sheet_view.showGridLines = False

    for col, w in zip("ABCDEFGHIJ", (18, 15, 15, 15, 15, 15, 15, 15, 15, 15)):
        ws.column_dimensions[col].width = w

    cur_month = wb[SET_MON]["B1"].value
    first_emp = None
    for r in range(EMP_ROWS[0], EMP_ROWS[1] + 1):
        if wb[SET_EMP].cell(r, 2).value:
            first_emp = wb[SET_EMP].cell(r, 2).value
            break

    put(ws, "A1", u"▶ 넣을 연월", 12, True, RED)
    put(ws, "B1", cur_month, 12, True, fill=FILL_IN, box=True, align="center")
    put(ws, "C1", u"← 이 두 칸만 고르면 됩니다", 10, color=RED)
    put(ws, "A2", u"▶ 넣을 사람", 12, True, RED)
    put(ws, "B2", first_emp, 12, True, fill=FILL_IN, box=True, align="center")

    dv_m, dv_e = _dv(u"연월목록"), _dv(u"직원목록")
    ws.add_data_validation(dv_m); dv_m.add(ws["B1"])
    ws.add_data_validation(dv_e); dv_e.add(ws["B2"])

    put(ws, "A4", u"■ 여기에 붙여넣으세요", 11, True)
    rows = [
        ("A5", u"붙여넣을 칸", "B5",
         u'="A"&(COUNTA(%s!$A$2:$A$4001)+2)' % RAW, 14, True),
        ("A6", u"이 달은 며칠", "B6",
         u'=IFERROR(DAY(EOMONTH(INDEX(%s!$B$4:$B$27,MATCH($B$1,%s!$A$4:$A$27,0)),0)),"")'
         % (SET_MON, SET_MON), 11, False),
        ("A7", u"지금 들어온 줄 수", "B7",
         u'=COUNTIF(%s!$V$2:$V$4001,$B$1&"|"&$B$2)' % RAW, 11, False),
    ]
    for la, lt, va, vf, sz, bold in rows:
        put(ws, la, lt, 11, True)
        put(ws, va, vf, sz, bold, fill=FILL_AUTO, box=True, align="center")

    # 판정 — 헛경고를 내지 않는 것이 중요합니다.
    #   · 고정월급(대표)은 근태를 아예 안 넣습니다 → 채근하지 않습니다
    #   · 중도 입·퇴사면 줄 수가 적은 것이 정상입니다 → 그렇게 말해 줍니다
    way = u'IFERROR(INDEX(%s!$I$2:$I$9,MATCH($B$2,%s,0)),"")' % (SET_EMP, emp_list_ref())
    put(ws, "A8", u"판정", 11, True)
    put(ws, "B8",
        u'=IF($B$2="","사람을 고르세요",'
        u'IF(%s="고정월급","고정월급이라 근태를 넣지 않습니다",'
        u'IF($B$7=0,"아직 안 넣었습니다 — 위에 적힌 칸에 붙여넣으세요",'
        u'IF($B$7=$B$6,"맞습니다",'
        u'IF($B$7=$B$6*2,"두 번 들어간 것 같습니다 — 겹친 줄을 지우세요",'
        u'IF($B$7<$B$6,$B$7&"줄입니다. 이 달("&$B$6&"일)보다 적습니다 — 중간에 들어오거나 나갔으면 정상입니다",'
        u'$B$7&"줄입니다. 이 달("&$B$6&"일)보다 많습니다 — 겹쳐 넣었는지 확인하세요'
        u'"))))))' % way,
        11, True, fill=FILL_AUTO, box=True)
    ws.merge_cells("B8:G8")

    put(ws, "I1",
        u"근태기록 .xls 를 열어 2행부터 마지막 데이터 줄까지 고르고(머리글과 합계 줄은 빼고) "
        u"복사한 뒤, 위에 적힌 칸을 눌러 붙여넣으세요. "
        u"사람마다 한 번씩, 다섯 사람이면 다섯 번입니다.",
        10, fill=FILL_NOTE, box=True, wrap=True)
    ws.merge_cells("I1:J8")

    put(ws, "A10", u"■ 근태원본 전체 점검", 11, True)
    for la, lt, va, vf in [
        ("A11", u"마지막 데이터 행", "B11",
         u'=IFERROR(LOOKUP(2,1/(%s!$A$2:$A$4001<>""),ROW(%s!$A$2:$A$4001)),1)' % (RAW, RAW)),
        ("A12", u"전체 줄 수", "B12", u'=COUNTA(%s!$A$2:$A$4001)' % RAW),
        ("A13", u"아래 표 합계", "B13", u"=SUM($B$18:$I$41)"),
    ]:
        put(ws, la, lt, 11, True)
        put(ws, va, vf, 11, False, fill=FILL_AUTO, box=True, align="center")

    put(ws, "A14", u"판정", 11, True)
    put(ws, "B14",
        u'=IF($B$11-1<>$B$12,"근태원본 중간에 빈 줄이 있습니다 (마지막 행 "&$B$11&", 줄 수 "&$B$12&")",'
        u'IF($B$13<>$B$12,"직원설정에 없는 이름이나 목록에 없는 달이 "&($B$12-$B$13)&"줄 섞여 있습니다",'
        u'"이상 없습니다"))',
        11, True, fill=FILL_AUTO, box=True)
    ws.merge_cells("B14:F14")

    put(ws, "A16", u"■ 넣은 것 한눈에 보기   (가로 사람 · 세로 연월 · 숫자는 줄 수)", 11, True)
    put(ws, "A17", u"연월", 9, True, INK_ON_DARK, FILL_HEAD, True, "center")
    for i in range(8):
        col = chr(ord("B") + i)
        put(ws, u"%s17" % col,
            u'=IF(%s!$B$%d="","",%s!$B$%d)' % (SET_EMP, EMP_ROWS[0] + i, SET_EMP, EMP_ROWS[0] + i),
            9, True, INK_ON_DARK, FILL_HEAD, True, "center")
    for j in range(24):
        r = 18 + j
        put(ws, u"A%d" % r, u"=%s!$A$%d" % (SET_MON, MON_ROWS[0] + j),
            10, False, box=True, align="center")
        for i in range(8):
            col = chr(ord("B") + i)
            put(ws, u"%s%d" % (col, r),
                u'=IF(%s$17="","",COUNTIF(%s!$V$2:$V$4001,$A%d&"|"&%s$17))'
                % (col, RAW, r, col),
                10, False, box=True, align="center")
    ws.freeze_panes = "B18"
    report.append(u"[%s] 시트를 만들었습니다 (근태원본 앞에)" % NEW_SHEET)
    return ws


# ─────────────────────────────────────────────────────────────
def run(path, out_path, do_write):
    wb = openpyxl.load_workbook(path, data_only=False)
    report = []

    gt = [n for n in wb.sheetnames if n.startswith(u"근퇴-")]
    print(u"■ 지금 시트 %d장" % len(wb.sheetnames))
    print(u"   근퇴-이름 : %d장  %s" % (len(gt), u", ".join(gt)))
    print(u"   %s 시트 : %s" % (NEW_SHEET, u"이미 있음" if NEW_SHEET in wb.sheetnames else u"없음"))
    print(u"\n■ 바꾼 뒤")
    print(u"   근퇴-이름 %d장 → [%s] 한 장 (이름 칸에서 사람을 고름)" % (len(gt), GT_NEW))
    print(u"   [%s] 시트 새로 추가 (연월·사람을 골라 넣을 자리를 안내하고 검사)" % NEW_SHEET)
    print(u"   시트 수 %d → %d" % (len(wb.sheetnames), len(wb.sheetnames) - len(gt) + 1 + 1))
    print(u"   근태원본은 **그대로 둡니다** (모든 달을 쌓아 두는 창고이므로)")

    if not do_write:
        print(u"\n※ 보기만 했습니다. 실제로 만들려면 --write 를 붙이세요.")
        return

    for nm, ref in ((u"직원목록", emp_list_ref()), (u"연월목록", mon_list_ref())):
        if nm in wb.defined_names:
            del wb.defined_names[nm]
        wb.defined_names[nm] = openpyxl.workbook.defined_name.DefinedName(nm, attr_text=ref)

    merge_gt(wb, report)
    make_input_sheet(wb, report)

    g = wb[GUIDE]
    for row in g.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith(u"근퇴-이름"):
                c.value = (u"근퇴계      사람을 골라 보는 근퇴계 + 급여산출 근거, 조회월 기준   [자동]")
            if isinstance(c.value, str) and c.value.startswith(u"근태원본    "):
                c.value = (u"근태원본    근태기록 원본을 그대로 붙여넣는 곳 (누적)           [입력]\n"
                           u"근태넣기    연월·사람을 골라 붙여넣을 자리를 안내하고 검사      [입력]")

    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    print(u"")
    for line in report:
        print(u"   " + line)
    print(u"\n만들었습니다: %s" % out_path)
    print(u"원본은 그대로 두었습니다: %s" % path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("book")
    ap.add_argument("--out", default=None)
    ap.add_argument("--write", action="store_true")
    a = ap.parse_args()
    out = a.out or re.sub(r"\.xlsx$", "", a.book) + u"_시트정리.xlsx"
    run(a.book, out, a.write)
