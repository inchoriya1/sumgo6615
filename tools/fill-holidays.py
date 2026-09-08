# -*- coding: utf-8 -*-
"""
[월설정] G열 공휴일 표를 채웁니다.

  python tools/fill-holidays.py 급여대장.xlsx            # 보기만
  python tools/fill-holidays.py 급여대장.xlsx --write    # 새 파일 만들기

왜 중요한가
  근무일명칭(평일/휴일/공휴일)이 이 표로 정해지고,
  그것이 수당 종류(연장 / 토요 / 특근)를 가릅니다.
  **표에 없는 공휴일은 평일로 잡히고 오류도 안 뜹니다.**

날짜를 어떻게 정했나
  · 날짜가 고정인 것(신정·삼일절·어린이날·현충일·광복절·개천절·한글날·성탄절)은 그대로.
  · 음력인 것(설날·부처님오신날·추석)은 2026년 양력 날짜를 확인해 넣었습니다.
  · 대체공휴일은 2026년에 **네 번**입니다 — 3/2 · 5/25 · 8/17 · 10/5.
    (삼일절·부처님오신날이 일요일, 광복절·개천절이 토요일과 겹쳤습니다.)
  · **8/17 대체공휴일은 이 회사 근태기록에 실제로 찍혀 있어서**,
    규칙이 맞다는 것이 자료로 확인되었습니다.

회사 휴일도 넣습니다
  **제헌절(7/17)** 은 나라의 공휴일이 아닙니다(2008년부터 제외).
  그런데 이 회사는 그날을 쉬었고 근태기록에 '제헌절'로 찍혀 있습니다.
  그래서 **회사 휴일**로 함께 넣었습니다.
  해마다 쉬는지는 회사가 정할 일이니, 2027년 것은 확인 후 넣으세요.
"""
import argparse
import datetime
import re

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SET_MON = u"월설정"
HOL_COL, HOL_ROW0, HOL_N = 7, 4, 60
FONT = u"맑은 고딕"
FILL_IN = PatternFill("solid", fgColor=u"FFFFE699")
THIN = Side(style="thin", color=u"FFD9D9D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def d(s):
    return datetime.date(*[int(x) for x in s.split("-")])

HOLIDAYS_2026 = [
    (d("2026-01-01"), u"신정"),
    (d("2026-02-16"), u"설날연휴"),
    (d("2026-02-17"), u"설날"),
    (d("2026-02-18"), u"설날연휴"),
    (d("2026-03-01"), u"삼일절"),
    (d("2026-03-02"), u"대체공휴일"),      # 삼일절이 일요일
    (d("2026-05-05"), u"어린이날"),
    (d("2026-05-24"), u"부처님오신날"),
    (d("2026-05-25"), u"대체공휴일"),      # 부처님오신날이 일요일
    (d("2026-06-03"), u"지방선거일"),      # 근태기록으로 확인됨
    (d("2026-06-06"), u"현충일"),          # 근태기록으로 확인됨
    (d("2026-07-17"), u"제헌절"),          # 회사 휴일 — 근태기록으로 확인됨
    (d("2026-08-15"), u"광복절"),          # 근태기록으로 확인됨
    (d("2026-08-17"), u"대체공휴일"),      # 광복절이 토요일 — 근태기록으로 확인됨
    (d("2026-09-24"), u"추석연휴"),
    (d("2026-09-25"), u"추석"),
    (d("2026-09-26"), u"추석연휴"),
    (d("2026-10-03"), u"개천절"),
    (d("2026-10-05"), u"대체공휴일"),      # 개천절이 토요일
    (d("2026-10-09"), u"한글날"),
    (d("2026-12-25"), u"성탄절"),
]

# 근태기록 원본에 실제로 찍혀 있던 날 — 여기와 어긋나면 안 됩니다.
FROM_DATA = {d("2026-06-03"), d("2026-06-06"), d("2026-07-17"),
             d("2026-08-15"), d("2026-08-17")}


def run(path, out_path, do_write):
    wb = openpyxl.load_workbook(path)
    if SET_MON not in wb.sheetnames:
        raise SystemExit(u"[%s] 시트가 없습니다." % SET_MON)
    ms = wb[SET_MON]

    now = []
    for i in range(HOL_N):
        v = ms.cell(HOL_ROW0 + i, HOL_COL).value
        if v is not None:
            now.append((v.date() if hasattr(v, "date") else v,
                        ms.cell(HOL_ROW0 + i, HOL_COL + 1).value))
    print(u"■ 지금 표에 %d개" % len(now))
    for v, nm in now:
        print(u"     %s  %s" % (v, nm))

    have = {v for v, _ in now}
    new = [(v, nm) for v, nm in HOLIDAYS_2026 if v not in have]
    lost = [(v, nm) for v, nm in now if v not in {x for x, _ in HOLIDAYS_2026}]
    print(u"\n■ 새로 넣을 것 %d개" % len(new))
    W = u"월화수목금토일"
    for v, nm in new:
        print(u"     %s (%s)  %-8s %s" % (v, W[v.weekday()], nm,
              u"← 근태기록으로 확인됨" if v in FROM_DATA else u""))
    if lost:
        print(u"\n■ 지금 표에 있는데 새 목록에 없는 것 (그대로 둡니다)")
        for v, nm in lost:
            print(u"     %s  %s" % (v, nm))

    # 6~8월은 이미 자료가 들어 있는 달입니다. 여기에 새 공휴일이 끼면
    # 지난달 급여가 바뀌므로, 그런 일이 없는지 확인합니다.
    risky = [(v, nm) for v, nm in new if 6 <= v.month <= 8]
    if risky:
        print(u"\n★ 6~8월에 새로 들어가는 공휴일이 있습니다 — 지난 급여가 바뀝니다!")
        for v, nm in risky:
            print(u"     %s %s" % (v, nm))
    else:
        print(u"\n■ 6~8월에 새로 들어가는 것 없음 — 지난 급여는 그대로입니다.")

    if not do_write:
        print(u"\n※ 보기만 했습니다. 실제로 만들려면 --write 를 붙이세요.")
        return

    allrows = sorted(set(now) | set(HOLIDAYS_2026))
    if len(allrows) > HOL_N:
        raise SystemExit(u"표 자리가 %d개인데 %d개를 넣으려 합니다." % (HOL_N, len(allrows)))
    for i in range(HOL_N):
        r = HOL_ROW0 + i
        for c in (HOL_COL, HOL_COL + 1):
            cell = ms.cell(r, c)
            cell.value = None
            cell.fill = FILL_IN
            cell.border = BOX
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(horizontal="center")
        ms.cell(r, HOL_COL).number_format = "yyyy-mm-dd"
    for i, (v, nm) in enumerate(allrows):
        ms.cell(HOL_ROW0 + i, HOL_COL).value = v
        ms.cell(HOL_ROW0 + i, HOL_COL + 1).value = nm

    ms["I3"] = (u"← 2026년은 다 채워져 있습니다. "
                u"해가 바뀌면 그해 것을 넣으세요. 비면 그날이 '평일'로 잡힙니다")
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    print(u"\n%d개를 날짜순으로 넣었습니다." % len(allrows))
    print(u"만들었습니다: %s" % out_path)
    print(u"원본은 그대로 두었습니다: %s" % path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("book")
    ap.add_argument("--out", default=None)
    ap.add_argument("--write", action="store_true")
    a = ap.parse_args()
    out = a.out or re.sub(r"\.xlsx$", "", a.book) + u"_공휴일.xlsx"
    run(a.book, out, a.write)
