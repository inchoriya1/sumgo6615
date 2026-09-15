# -*- coding: utf-8 -*-
u"""
시트 세 곳을 정리합니다 — 안내문·근태점검·조회월 표시.

  python tools/tidy-sheets.py 급여대장.xlsx            # 보기만
  python tools/tidy-sheets.py 급여대장.xlsx --write    # 새 파일 만들기
  python tools/tidy-sheets.py 급여대장.xlsx --write --out 같은이름.xlsx   # 덮어쓰기

무엇이 문제였나
  ① **[사용안내] 가 옛날 방식을 그대로 적고 있습니다.**
     "맨 아래 빈 행에 붙여넣기 · 계속 쌓아 두면" 이라고 적혀 있는데,
     지금 [근태원본] 은 24개월 × 8명 자리가 미리 깔린 고정판이라
     필터로 골라 노란 칸에 넣는 방식입니다. 안내와 시트가 어긋나 있었습니다.
     곁들여, 시트 목록의 '근태원본'과 '근태넣기'가 **한 칸에 두 줄로 들어가 있어
     둘째 줄이 화면에 안 보였습니다.** 줄을 나눠 둘 다 보이게 합니다.

  ② **[근태넣기] 라는 이름이 넣는 곳처럼 보입니다.**
     실제로는 아무 시트도 이 시트를 참조하지 않습니다. 계산에 전혀 안 쓰이는
     **순수 점검판**입니다. 이름을 [근태점검] 으로 바꿉니다.
     위쪽 '이 달 공휴일 몇 개' 줄은 [월설정] G열을 직접 보면 되는 내용이라 뺍니다.
     쓸모 있는 것은 아래 매트릭스 하나뿐이고, 그건 그대로 둡니다 —
     5,952줄에 필터로 밀어 넣다 보면 '8월 그 사람 넣었던가'를
     눈으로 확인할 방법이 사실상 그것뿐입니다.

  ③ **필터를 두 번 걸어야 합니다 — 연월 하나, 이름 하나.**
     [근태원본] Y열에 조회월 열을 만듭니다. [월설정] B1 을 바꾸면
     그 달 248줄에만 표시가 켜집니다. Y열 필터 하나만 걸면 한 달이 통째로 나옵니다.

     **근무일자는 여전히 조회월 따라 움직이지 않습니다. 그래야 맞습니다.**
     출근·퇴근은 수식이 아니라 직접 친 값이라, 날짜가 조회월을 따라 움직이면
     8월에 넣어 둔 출퇴근이 9월로 바꾸는 순간 9월 근태로 둔갑합니다.
     Y열은 날짜를 옮기지 않고 '찾아가는 길'만 만들어 줍니다.

  ④ 곁들여, **조건부서식이 4001행에서 끊겨 있었습니다.**
     사람 블록마다 번갈아 칠하는 줄무늬와 경계선이 2027년 초부터 사라집니다.
     4000행짜리이던 시절의 자취입니다. 5953행까지, Y열까지 늘립니다.
"""
import argparse
import re

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAW = u"근태원본"
OLD_CHECK = u"근태넣기"
NEW_CHECK = u"근태점검"
GUIDE = u"사용안내"

FIRST, LAST = 2, 5953
MONTH_ROWS = 24                 # 월설정 A4:A27
SLOTS = 8                       # 직원설정 B2:B9
FLAG_COL = 25                   # Y — 조회월 표시
NOTE_COL = 26                   # Z — 넣는 법 안내
MARK = u"◀"

FONT = u"맑은 고딕"
INK_HEAD = u"FF1F4E79"          # 진한 남색 — 표 머리글
INK_HELP = u"FF808080"          # 회색 — 파생(자동) 열 머리글
INK_NOTE = u"FFFFF9E6"          # 연노랑 — 안내 상자
THIN = Border(*(Side(style="thin"),) * 4)


# ── ① 사용안내 ────────────────────────────────────────────────────────
GUIDE_TEXT = [
    (5,  u"① [근태원본]  1행 필터에서 연월과 이름을 고르고, 노란 칸(출근·퇴근)에 넣기"),
    (6,  u"     · 자리는 24개월 × 8명이 미리 깔려 있습니다. 줄을 더하거나 지우지 마세요."),
    (7,  u"     · 같은 자리에 다시 넣으면 덮어써집니다. 두 배가 되지 않습니다."),
    (13, u"     · [근태원본] 의 근무일자는 조회월을 따라가지 않습니다. 24개월치가 늘 깔려 있습니다."),
    (15, u"월설정      조회월 · 회사명 · 월별 임금산정기간과 지급일        [입력]"),
    (16, u"근태원본    연월·이름을 필터로 골라 출퇴근을 넣는 곳 (24개월 고정판)   [입력]"),
    (17, u"근태점검    어느 달 누구를 넣었는지 한눈에 보는 곳 (계산에는 안 쓰임)  [확인]"),
    (18, u"직원설정    직원의 변하지 않는 정보와 기본값                    [입력]"),
    (19, u"월별입력    연월 × 직원별로 달라지는 값(시급·기본근로시간·유예·보험 적용 등)   [입력]"),
    (20, u"근태집계    근태원본에서 연월 × 직원별 자동 집계                [자동]"),
    (21, u"기본정보    급여 계산 엔진. 모든 달 · 모든 직원을 계산          [자동]"),
    (22, u"전체임금대장  조회월 기준 임금대장                          [자동]"),
    (23, u"신-급여명세서 · 명세서교부대장               조회월 기준         [자동]"),
    (24, u"근퇴계      사람을 골라 보는 근퇴계 + 급여산출 근거, 조회월 기준   [자동]"),
    (25, u"연간집계    연도별 지급 · 공제 · 차인지급액 매트릭스            [자동]"),
    (26, u""),
]


def fix_guide(wb):
    ws = wb[GUIDE]
    for row, text in GUIDE_TEXT:
        c = ws.cell(row=row, column=2)
        c.value = text or None
        c.font = Font(name=FONT, size=11)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    return len(GUIDE_TEXT)


# ── ② 근태점검 ────────────────────────────────────────────────────────
CHECK_HOWTO = u"""이 시트는 넣는 곳이 아니라, 다 넣었는지 확인하는 곳입니다.
계산에는 쓰이지 않습니다 — 지워도 급여는 한 원도 안 바뀝니다.

넣는 곳은 [근태원본] 입니다.
① [근태원본] 시트로 갑니다.
② Y열 [조회월] 필터에서 %s 를 고르면 — [월설정] B1 의 달 248줄이 통째로 나옵니다.
   한 사람만 보려면 A열 [이름] 필터를 더 겁니다. 다른 달은 Q열 [연월] 필터로 고릅니다.
③ 노란 칸(출근·퇴근)에 치거나, 근태기록 .xls 의 출근·퇴근 두 열을 붙여넣습니다.
④ 돌아와 아래 표를 봅니다. 빈칸이면 안 들어간 것, 숫자면 그만큼 들어온 것입니다.""" % MARK


def rebuild_check(wb):
    idx = wb.sheetnames.index(OLD_CHECK)
    del wb[OLD_CHECK]
    ws = wb.create_sheet(NEW_CHECK, idx)
    ws.sheet_view.showGridLines = False

    ws.column_dimensions["A"].width = 16.0
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 13.0

    t = ws["A1"]
    t.value = NEW_CHECK
    t.font = Font(name=FONT, size=16, bold=True)

    ws.merge_cells("A2:H8")
    n = ws["A2"]
    n.value = CHECK_HOWTO
    n.font = Font(name=FONT, size=11)
    n.fill = PatternFill("solid", fgColor=INK_NOTE)
    n.alignment = Alignment(vertical="center", wrap_text=True)
    n.border = THIN
    ws.row_dimensions[2].height = 110.0

    h = ws["A10"]
    h.value = u"■ 넣은 것 한눈에 보기"
    h.font = Font(name=FONT, size=11, bold=True)
    ws.merge_cells("B10:H10")
    s = ws["B10"]
    s.value = u"가로 사람 · 세로 연월 · 숫자는 출퇴근이 채워진 날 수"
    s.font = Font(name=FONT, size=10, color=INK_HELP)

    head = Font(name=FONT, size=9, bold=True, color="FFFFFFFF")
    fill = PatternFill("solid", fgColor=INK_HEAD)
    mid = Alignment(horizontal="center", vertical="center")

    HR = 11                                  # 표 머리글 행
    ws.cell(row=HR, column=1, value=u"연월")
    for s_i in range(SLOTS):                 # B..I ← 직원설정 B2..B9
        ws.cell(row=HR, column=2 + s_i,
                value=u'=IF(직원설정!$B$%d="","",직원설정!$B$%d)' % (2 + s_i, 2 + s_i))
    for c in ws[HR][:1 + SLOTS]:
        c.font, c.fill, c.alignment, c.border = head, fill, mid, THIN

    body = Font(name=FONT, size=10)
    for m_i in range(MONTH_ROWS):            # 24행 ← 월설정 A4..A27
        r = HR + 1 + m_i
        a = ws.cell(row=r, column=1, value=u"=월설정!$A$%d" % (4 + m_i))
        a.font, a.alignment, a.border = body, mid, THIN
        for s_i in range(SLOTS):
            col = 2 + s_i
            L = get_column_letter(col)
            c = ws.cell(row=r, column=col, value=(
                u'=IF(%s$%d="","",COUNTIFS('
                u'근태원본!$V$%d:$V$%d,$A%d&"|"&%s$%d,'
                u'근태원본!$D$%d:$D$%d,"<>"))'
                % (L, HR, FIRST, LAST, r, L, HR, FIRST, LAST)))
            c.font, c.alignment, c.border = body, mid, THIN

    ws.freeze_panes = "B%d" % (HR + 1)
    return MONTH_ROWS * SLOTS


# ── ③ 근태원본 Y열 ────────────────────────────────────────────────────
RAW_NOTE = u"""■ 매월 이렇게 넣습니다
① Y열 [조회월] 필터에서 %s 를 고릅니다.
   → [월설정] B1 의 달, 여덟 사람 248줄이 통째로 남습니다.
   한 사람만 보려면 A열 [이름] 필터를 더 겁니다.
   다른 달은 Q열 [연월] 필터에서 직접 고릅니다.
② 노란 칸(출근·퇴근) 에 치거나, 근태기록 .xls 의
   출근·퇴근 두 열을 복사해 그대로 붙여넣습니다.
③ 다 넣었는지는 [근태점검] 시트에서 한눈에 봅니다.

· 자리는 미리 깔려 있습니다. 줄을 더하거나 지우지 마세요.
· 같은 자리에 다시 넣으면 덮어써집니다. 두 배가 되지 않습니다.
· 근무일자는 조회월을 따라 움직이지 않습니다. 24개월치가 늘
  깔려 있어서입니다. 그래야 지난 달에 넣어 둔 출퇴근이 그 달에 남습니다.
· 이름·날짜·요일·근무일명칭·기본·연장은 저절로 채워집니다.
· 지각을 깎을 때만 O열 [지각차감(H)] 에 시간을 적습니다.
· 공휴일은 [월설정] G열 표를 봅니다. 비면 평일로 잡힙니다.
· 전체를 다시 보려면 필터에서 '모두 선택' 을 고르세요.""" % MARK


def add_flag_column(wb):
    ws = wb[RAW]
    L = get_column_letter(FLAG_COL)

    h = ws.cell(row=1, column=FLAG_COL, value=u"조회월")
    h.font = Font(name=FONT, size=9, bold=True, color="FFFFFFFF")
    h.fill = PatternFill("solid", fgColor=INK_HELP)
    h.alignment = Alignment(horizontal="center", vertical="center")
    h.border = THIN
    ws.column_dimensions[L].width = 10.0

    body = Font(name=FONT, size=11)
    mid = Alignment(horizontal="center", vertical="center")
    for r in range(FIRST, LAST + 1):
        c = ws.cell(row=r, column=FLAG_COL,
                    value=u'=IF($Q%d="","",IF($Q%d=조회월,"%s",""))' % (r, r, MARK))
        c.font, c.alignment, c.border = body, mid, THIN

    ws.cell(row=1, column=NOTE_COL).value = RAW_NOTE

    # 필터와 조건부서식을 Y열 · 5953행까지 늘립니다
    ref = u"A1:%s%d" % (L, LAST)
    ws.auto_filter.ref = ref
    rules = [rule for cf in ws.conditional_formatting for rule in cf.rules]
    ws.conditional_formatting = ConditionalFormattingList()
    for rule in rules:
        ws.conditional_formatting.add(u"A%d:%s%d" % (FIRST, L, LAST), rule)
    return ref, len(rules)


def run(path, out_path, do_write):
    wb = openpyxl.load_workbook(path)
    old_filter = wb[RAW].auto_filter.ref
    old_cf = [str(cf.sqref) for cf in wb[RAW].conditional_formatting]

    n_guide = fix_guide(wb)
    n_cells = rebuild_check(wb)
    new_filter, n_rules = add_flag_column(wb)

    print(u"① [사용안내]  %d 줄을 지금 방식으로 고쳤습니다" % n_guide)
    print(u"              시트 목록의 근태원본·근태넣기 두 줄 겹침도 풀었습니다")
    print(u"② [근태넣기] → [%s]  공휴일 줄을 빼고 매트릭스(%d칸)만 남겼습니다"
          % (NEW_CHECK, n_cells))
    print(u"③ [근태원본]  Y열 조회월 추가 — %d줄" % (LAST - FIRST + 1))
    print(u"              필터   %s → %s" % (old_filter, new_filter))
    print(u"              조건부서식 %s → A2:%s%d  (규칙 %d개)"
          % (", ".join(old_cf), L_OF_FLAG, LAST, n_rules))

    if not do_write:
        print(u"\n※ 보기만 했습니다. 실제로 만들려면 --write 를 붙이세요.")
        return
    wb.save(out_path)
    print(u"\n만들었습니다: %s" % out_path)


L_OF_FLAG = get_column_letter(FLAG_COL)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("book")
    ap.add_argument("--out", default=None)
    ap.add_argument("--write", action="store_true")
    a = ap.parse_args()
    out = a.out or re.sub(r"\.xlsx$", "", a.book) + u"_정리.xlsx"
    run(a.book, out, a.write)


if __name__ == "__main__":
    main()
