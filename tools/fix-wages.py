# -*- coding: utf-8 -*-
u"""
시급을 회사 근퇴계와 맞추고, 기본정보의 죽은 중복 열을 걷어냅니다.

  python tools/fix-wages.py 급여대장.xlsx            # 보기만
  python tools/fix-wages.py 급여대장.xlsx --write    # 새 파일 만들기
  python tools/fix-wages.py 급여대장.xlsx --write --out 같은이름.xlsx   # 덮어쓰기

① 시급이 틀려 있었습니다 — 근퇴계가 안 맞던 진짜 이유

  회사에서 받은 8월 근퇴계의 시급·일급과 대조해 보니 두 사람이 어긋났습니다.

      209H 두 사람    받은 값과 이미 일치 — 그대로
      일급제 한 사람   ★ 시급이 소수점까지 들어가 일급이 어긋남
      일수 한 사람     ★ 시급 자체가 달랐음 (가장 큰 차이)

  **일급 = 시급 × 8** 이 정확히 맞는 값을 역산해 넣습니다.
  · 일급제 → 받은 일급 ÷ 8 (정수로 딱 떨어짐)
  · 일수  → 받은 일급 ÷ 8 (소수점 세 자리까지)

  받은 근퇴계의 시급은 **화면에 반올림돼 보이는 값**입니다.
  그 값을 그대로 넣으면 일급이 몇 원 틀어집니다.
  그래서 **일급 ÷ 8** 로 역산한 값을 넣습니다.

  자료를 못 받은 한 사람은 **손대지 않았습니다.**

② 통상일급이 10원 단위로 잘려 있었습니다

      전  일급제 → ROUND(시급×8, -1)      10원이 잘림
      후  일급제 → 시급×8                  받은 근퇴계와 일치

  209H·일수는 원래도 `시급×8` 이었고, 고정월급만 원 단위 반올림입니다.
  **일급제만 혼자 10원 절사**를 하고 있어서 받은 근퇴계와 어긋났습니다.

③ 기본정보에 아무도 안 보는 열이 셋 있었습니다

  192행짜리 열이 **참조 0곳**으로 놀고 있었습니다.

      C   주민번호   ← 아무도 안 봄. **주민등록번호가 192줄**
      H   비고       ← 아무도 안 봄
      BA  비고       ← H 와 똑같은 수식. 완전 중복

  세 열을 비웁니다. A·B·D~G 는 전체임금대장·급여명세서가 실제로
  가져다 쓰므로 **그대로 둡니다** — 중복처럼 보여도 그게 조회표 역할입니다.

  **특히 C** 는 지울 값어치가 큽니다. 쓰지도 않는 곳에 주민등록번호가
  192번 복제돼 있었습니다. 원본은 [직원설정] 에 그대로 있습니다.

④ 근퇴계에 통상시급·통상일급을 띄웁니다

  받은 근퇴계와 **대조할 값이 화면에 없었습니다.** 급여산출 블록 맨 위에
  두 줄을 넣어, 숫자가 어긋나면 **여기부터 보게** 합니다.
"""
import argparse
import re
from copy import copy

import openpyxl
from openpyxl.styles import Alignment, Font

EMP, BASE, GT = u"직원설정", u"기본정보", u"근퇴계"

# 직원설정 L(기본시급) 정정 — 받은 근퇴계의 일급 ÷ 8
WAGE_FIX = {
    # 받은 근퇴계의 일급 ÷ 8. 이름과 금액은 저장소에 적지 않습니다 —
    # 쓰실 때 이 두 줄만 실제 이름·일급으로 채워 돌리세요.
    # u"<이름>": <받은 일급> / 8.0,
}
# 확인만 하고 그대로 두는 사람 (받은 자료와 이미 일치)
WAGE_OK = {}      # 이미 맞는 사람 (확인만 하고 그대로 둠)

BASE_ROWS = (2, 193)
DEAD_COLS = ((3, u"주민번호"), (8, u"비고"), (53, u"비고"))   # C · H · BA

DAY_OLD = u'=IF(직원설정!I2="일급제",ROUND($I2*8,-1),IF(직원설정!I2="고정월급",ROUND($I2*8,0),$I2*8))'
DAY_NEW = u'=IF(직원설정!I{r}="고정월급",ROUND($I{r}*8,0),$I{r}*8)'

FONT = u"맑은 고딕"


def fix_wages(wb):
    ws = wb[EMP]
    out = []
    for r in range(2, 10):
        nm = ws.cell(r, 2).value
        if not nm:
            continue
        cell = ws.cell(r, 12)                      # L 기본시급
        if nm in WAGE_FIX:
            out.append((nm, cell.value, WAGE_FIX[nm], True))
            cell.value = WAGE_FIX[nm]
        elif nm in WAGE_OK:
            ok = abs((cell.value or 0) - WAGE_OK[nm]) < 0.005
            out.append((nm, cell.value, WAGE_OK[nm], not ok))
            if not ok:
                cell.value = WAGE_OK[nm]
    return out


def fix_daily(wb):
    u"""통상일급(K) — 일급제의 10원 절사를 없애 '시급 × 8' 로 맞춥니다."""
    ws = wb[BASE]
    first, last = BASE_ROWS
    before = ws.cell(first, 11).value
    for r in range(first, last + 1):
        ws.cell(row=r, column=11).value = DAY_NEW.format(r=r)
    return before, ws.cell(first, 11).value


def clear_dead(wb):
    u"""아무도 참조하지 않는 열을 비웁니다. 머리글에 이유를 적어 둡니다."""
    ws = wb[BASE]
    first, last = BASE_ROWS
    done = []
    for col, name in DEAD_COLS:
        n = 0
        for r in range(first, last + 1):
            c = ws.cell(row=r, column=col)
            if c.value is not None:
                c.value = None
                n += 1
        h = ws.cell(row=1, column=col)
        old = h.value
        h.value = u"(사용 안 함)"
        h.font = Font(name=FONT, size=9, italic=True, color="FF808080")
        done.append((openpyxl.utils.get_column_letter(col), old, name, n))
    return done


# 근퇴계 급여산출 블록은 M~P 7~26행. 그 위 5~6행에 두 줄을 끼웁니다.
GT_ROWS = [
    (5, u"통상시급", u'=INDEX(기본정보!$I:$I,$T$1)', u'#,##0.##;\\-#,##0.##;"-"'),
    (6, u"통상일급", u'=INDEX(기본정보!$K:$K,$T$1)', u'#,##0;\\-#,##0;"-"'),
]


def add_gt_wage(wb):
    ws = wb[GT]
    mid = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    src = ws["M8"]                                  # 블록 머리글 서식을 빌려 씁니다
    for row, label, formula, fmt in GT_ROWS:
        a = ws.cell(row=row, column=13, value=label)     # M
        a.font = copy(src.font)
        a.fill = copy(src.fill)
        a.border = copy(src.border)
        a.alignment = mid
        b = ws.cell(row=row, column=16, value=formula)    # P 금액 열
        b.font = Font(name=FONT, size=11, bold=True)
        b.border = copy(src.border)
        b.alignment = right
        b.number_format = fmt
    note = ws.cell(row=5, column=14, value=u"← 받은 근퇴계와 먼저 맞춰 보세요")
    note.font = Font(name=FONT, size=9, color="FF808080")
    return [r for r, _, _, _ in GT_ROWS]


def run(path, out_path, do_write):
    wb = openpyxl.load_workbook(path)

    wages = fix_wages(wb)
    d_before, d_after = fix_daily(wb)
    dead = clear_dead(wb)
    gt = add_gt_wage(wb)

    print(u"① 직원설정 기본시급 — 받은 근퇴계와 대조")
    for nm, old, new, changed in wages:
        print(u"   %-6s %14.4f → %-14.4f 일급 %s  %s"
              % (nm, old or 0, new, format(new * 8, ",.0f"),
                 u"★ 고쳤습니다" if changed else u"맞음(그대로)"))
    print(u"   자료 없는 사람 — 손대지 않았습니다")
    print(u"② 기본정보 통상일급(K) %d줄" % (BASE_ROWS[1] - BASE_ROWS[0] + 1))
    print(u"   전: %s" % d_before)
    print(u"   후: %s" % d_after)
    print(u"③ 기본정보 죽은 열 비움 (참조 0곳)")
    for letter, old, name, n in dead:
        print(u"   %-3s %-8s %d칸 비움 → 머리글 '(사용 안 함)'" % (letter, name, n))
    print(u"④ 근퇴계 %d·%d행에 통상시급·통상일급 추가" % (gt[0], gt[1]))

    if not do_write:
        print(u"\n※ 보기만 했습니다. 실제로 만들려면 --write 를 붙이세요.")
        return
    wb.save(out_path)
    print(u"\n만들었습니다: %s" % out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("book")
    ap.add_argument("--out", default=None)
    ap.add_argument("--write", action="store_true")
    a = ap.parse_args()
    out = a.out or re.sub(r"\.xlsx$", "", a.book) + u"_시급.xlsx"
    run(a.book, out, a.write)


if __name__ == "__main__":
    main()
