# -*- coding: utf-8 -*-
"""
근태원본을 '한 사람만 골라 보고 거기서 바로 넣는' 화면으로 만듭니다.

  python tools/attendance-picker.py 급여대장.xlsx            # 보기만
  python tools/attendance-picker.py 급여대장.xlsx --write    # 새 파일 만들기

  ※ tools/rebuild-attendance.py 를 먼저 돌린 파일에만 씁니다.

무엇이 문제였나
  자리를 미리 깔아 두니 5,952줄이 되었고, 열면 **여덟 사람 스물네 달이
  한꺼번에** 쏟아졌습니다. 넣을 자리를 알려 줘도 그 줄까지 굴러가야 했습니다.

무엇을 하나
  근태원본에 **자동 필터**를 걸고, **이름**과 **연월**로 걸러 놓습니다.
  그러면 그 사람 그 달 **서른한 줄만** 남습니다.
  거기 노란 칸(출근·퇴근)에 바로 치거나 붙여넣으면 됩니다.
  다른 사람으로 옮길 때는 **필터 화살표에서 이름만 바꿔** 고르면 됩니다.

  자리는 그대로 있고 값만 채워지므로 **누적은 그대로**입니다.
  줄을 더하지 않으니 두 번 넣어 두 배가 될 일도 없습니다.

왜 드롭다운 한 칸으로 못 하나
  칸을 골라 화면이 바뀌게 하려면 매크로가 있어야 합니다.
  이 파일은 매크로 없는 .xlsx 로 두는 편이 안전해서(보안 경고·차단이 없습니다)
  엑셀이 원래 가지고 있는 **필터**를 쓴 것입니다.
"""
import argparse
import re

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

RAW, SET_EMP, SET_MON, IN = u"근태원본", u"직원설정", u"월설정", u"근태넣기"
SLOTS, MONTHS, DAYS = 8, 24, 31
FIRST = 2
LAST = FIRST + SLOTS * MONTHS * DAYS - 1
EMP_ROW0, MON_ROW0 = 2, 4
COL_NAME, COL_YM = 0, 16          # 자동필터에서 A열 = 0, Q열 = 16

FONT = u"맑은 고딕"
RED = u"FFC00000"
FILL_IN = PatternFill("solid", fgColor=u"FFFFE699")
FILL_AUTO = PatternFill("solid", fgColor=u"FFF2F2F2")
FILL_HEAD = PatternFill("solid", fgColor=u"FF1F4E79")
FILL_NOTE = PatternFill("solid", fgColor=u"FFFFF9E6")
THIN = Side(style="thin", color=u"FFD9D9D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def put(ws, addr, value, size=11, bold=False, color=None, fill=None,
        box=False, align=None, wrap=False):
    c = ws[addr]
    c.value = value
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    if fill:
        c.fill = fill
    if box:
        c.border = BOX
    if align or wrap:
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    return c


def block_start(month_i, slot_i):
    return FIRST + (month_i * SLOTS + slot_i) * DAYS


def apply_filter(wb, report):
    """자동 필터를 걸고, 첫 사람 · 조회월만 남기고 접어 둡니다."""
    raw = wb[RAW]
    month = wb[SET_MON]["B1"].value
    months = [wb[SET_MON].cell(MON_ROW0 + i, 1).value for i in range(MONTHS)]
    names = [wb[SET_EMP].cell(EMP_ROW0 + i, 2).value for i in range(SLOTS)]
    mi = months.index(month) if month in months else 0
    # 처음 열었을 때 보일 사람은 **근태를 넣는 사람**으로 고릅니다.
    # 고정월급(대표)은 근태가 없어서, 그 사람이 걸려 있으면 빈 화면만 보입니다.
    ways = [wb[SET_EMP].cell(EMP_ROW0 + i, 9).value for i in range(SLOTS)]
    si = next((i for i, n in enumerate(names) if n and ways[i] != u"고정월급"),
              next((i for i, n in enumerate(names) if n), 0))
    person = names[si]

    raw.auto_filter.ref = u"A1:X%d" % LAST
    raw.auto_filter.filterColumn = []
    raw.auto_filter.add_filter_column(COL_NAME, [person])
    raw.auto_filter.add_filter_column(COL_YM, [month])

    top = block_start(mi, si)
    for r in range(FIRST, LAST + 1):
        raw.row_dimensions[r].hidden = not (top <= r < top + DAYS)

    raw["Z1"] = (
        u"■ 매월 이렇게 넣습니다\n"
        u"① 1행의 필터 화살표에서 **이름**과 **연월**을 고릅니다.\n"
        u"   → 그 사람 그 달 서른한 줄만 남습니다.\n"
        u"② **노란 칸(출근·퇴근)** 에 치거나, 근태기록 .xls 의\n"
        u"   출근·퇴근 두 열을 복사해 그대로 붙여넣습니다.\n"
        u"③ 다음 사람은 **이름 필터만** 바꿔 고르면 됩니다.\n\n"
        u"· 자리는 미리 깔려 있습니다. 줄을 더하거나 지우지 마세요.\n"
        u"· 같은 자리에 다시 넣으면 덮어써집니다. 두 배가 되지 않습니다.\n"
        u"· 이름·날짜·요일·근무일명칭·기본·연장은 저절로 채워집니다.\n"
        u"· 지각을 깎을 때만 O열 [지각차감(H)] 에 시간을 적습니다.\n"
        u"· 공휴일은 [월설정] G열 표를 봅니다. 비면 평일로 잡힙니다.\n"
        u"· 전체를 다시 보려면 필터에서 '모두 선택' 을 고르세요.")
    report.append(u"근태원본에 자동 필터를 걸고 [%s · %s] 만 남겼습니다 (%d~%d행)"
                  % (person, month, top, top + DAYS - 1))
    return person, month


def rebuild_board(wb, report):
    """근태넣기를 '고르는 곳' 이 아니라 '현황판' 으로 다시 만듭니다."""
    if IN in wb.sheetnames:
        del wb[IN]
    ws = wb.create_sheet(IN, wb.sheetnames.index(RAW))
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJ", (16, 13, 13, 13, 13, 13, 13, 13, 13, 13)):
        ws.column_dimensions[col].width = w

    put(ws, "A1", u"근태 넣는 법", 16, True)
    put(ws, "A2",
        u"① [근태원본] 시트로 갑니다.\n"
        u"② 1행의 필터 화살표에서 이름과 연월을 고릅니다 — 그 사람 그 달 31줄만 남습니다.\n"
        u"③ 노란 칸(출근·퇴근)에 치거나, 근태기록 .xls 의 출근·퇴근 두 열을 붙여넣습니다.\n"
        u"④ 다음 사람은 이름 필터만 바꿔 고릅니다.\n\n"
        u"이 시트는 넣는 곳이 아니라 **다 넣었는지 확인하는 곳**입니다.",
        11, fill=FILL_NOTE, box=True, wrap=True)
    ws.merge_cells("A2:H7")
    ws.row_dimensions[2].height = 96

    put(ws, "A9", u"■ 이 달 공휴일", 11, True)
    put(ws, "B9",
        u'=COUNTIFS(공휴일날짜,">="&기간시작,공휴일날짜,"<="&기간종료)',
        12, True, fill=FILL_AUTO, box=True, align="center")
    put(ws, "C9",
        u'=IF($B$9=0,"★ 조회월("&조회월&")의 공휴일이 하나도 등록돼 있지 않습니다. '
        u'[월설정] G열을 채우세요","개 등록됨 ("&조회월&") — 달력과 맞는지 확인하세요")',
        10, color=RED)
    ws.merge_cells("C9:H9")

    put(ws, "A11", u"■ 넣은 것 한눈에 보기", 11, True)
    put(ws, "B11", u"가로 사람 · 세로 연월 · 숫자는 출퇴근이 채워진 날 수", 10, color=u"FF808080")
    ws.merge_cells("B11:H11")
    put(ws, "A12", u"연월", 9, True, u"FFFFFFFF", FILL_HEAD, True, "center")
    for i in range(SLOTS):
        col = chr(ord("B") + i)
        put(ws, u"%s12" % col,
            u'=IF(%s!$B$%d="","",%s!$B$%d)' % (SET_EMP, EMP_ROW0 + i, SET_EMP, EMP_ROW0 + i),
            9, True, u"FFFFFFFF", FILL_HEAD, True, "center")
    for j in range(MONTHS):
        r = 13 + j
        put(ws, u"A%d" % r, u"=%s!$A$%d" % (SET_MON, MON_ROW0 + j), 10, box=True, align="center")
        for i in range(SLOTS):
            col = chr(ord("B") + i)
            put(ws, u"%s%d" % (col, r),
                u'=IF(%s$12="","",COUNTIFS(%s!$V$%d:$V$%d,$A%d&"|"&%s$12,'
                u'%s!$D$%d:$D$%d,"<>"))'
                % (col, RAW, FIRST, LAST, r, col, RAW, FIRST, LAST),
                10, box=True, align="center")
    ws.freeze_panes = "B13"
    report.append(u"[%s] 를 현황판으로 다시 만들었습니다 (고르는 칸 없음)" % IN)


def run(path, out_path, do_write):
    wb = openpyxl.load_workbook(path, data_only=False)
    raw = wb[RAW]
    if raw.max_row < LAST:
        raise SystemExit(u"먼저 tools/rebuild-attendance.py 를 돌리세요 "
                         u"(지금 %d행, 필요 %d행)" % (raw.max_row, LAST))
    print(u"■ 지금 — 근태원본 %d줄이 한꺼번에 보입니다 (여덟 사람 × 스물네 달)"
          % (LAST - FIRST + 1))
    print(u"■ 바꾼 뒤 — 자동 필터로 **한 사람 한 달 31줄만** 보입니다")
    print(u"   고르기 : 근태원본 1행의 이름·연월 필터 화살표")
    print(u"   넣기   : 그 자리 노란 칸(출근·퇴근)에 바로")
    print(u"   근태넣기는 '현황판' 이 됩니다 — 다 넣었는지 확인하는 곳")
    if not do_write:
        print(u"\n※ 보기만 했습니다. 실제로 만들려면 --write 를 붙이세요.")
        return
    report = []
    apply_filter(wb, report)
    rebuild_board(wb, report)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    print(u"")
    for line in report:
        print(u"   " + line)
    print(u"\n만들었습니다: %s" % out_path)
    print(u"원본은 그대로 두었습니다: %s" % path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("book")
    ap.add_argument("--out", default=None)
    ap.add_argument("--write", action="store_true")
    a = ap.parse_args()
    out = a.out or re.sub(r"\.xlsx$", "", a.book) + u"_고르기.xlsx"
    run(a.book, out, a.write)
